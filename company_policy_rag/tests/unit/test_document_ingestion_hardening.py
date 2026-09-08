from __future__ import annotations

import zipfile
from unittest.mock import MagicMock

import pytest

from backend.ingestion.loaders.loader_factory import load_document
from backend.ingestion.loaders.validation import validate_office_archive


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-32"])
@pytest.mark.parametrize("extension", ["txt", "md", "csv", "json"])
def test_unicode_documents(tmp_path, encoding, extension):
    path = tmp_path / f"policy.{extension}"
    text = '{"policy": "नमस्ते café"}' if extension == "json" else "नमस्ते café"
    path.write_bytes(text.encode(encoding))
    assert "नमस्ते café" in load_document(path)[0].content


@pytest.mark.parametrize(
    "name,content,message",
    [
        ("file.exe", b"MZ executable", "Unsupported"),
        ("file.doc", b"legacy", "Unsupported"),
        ("file.txt", b"\x00binary", "Binary"),
        ("file.txt", b" \n", "No readable"),
        ("file.json", b'{"broken":', "Invalid JSON"),
        ("file.jsonl", b'{"ok":1}\n\ninvalid\n', "line 3"),
        ("file.docx", b"not a zip", "Invalid Office"),
    ],
)
def test_bad_documents_are_explicit_failures(tmp_path, name, content, message):
    path = tmp_path / name
    path.write_bytes(content)
    with pytest.raises(ValueError, match=message):
        load_document(path)


def test_csv_retains_extra_columns_and_quoted_newlines(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text('Name\n"Alice\nSmith",extra|value\n', encoding="utf-8")
    document = load_document(path)[0]
    assert "Alice Smith" in document.content
    assert "extra&#124;value" in document.content
    assert document.metadata.extra["columns"] == ["Name", "Column 2"]


def test_tsv(tmp_path):
    path = tmp_path / "data.tsv"
    path.write_text("name\tamount\nAlice\t42", encoding="utf-8")
    assert "| Alice | 42 |" in load_document(path)[0].content


def test_office_expansion_limit(tmp_path, monkeypatch):
    from backend.ingestion.loaders import validation

    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 1000)
    monkeypatch.setattr(validation, "MAX_EXPANDED_BYTES", 100)
    with pytest.raises(ValueError, match="expanded-content"):
        validate_office_archive(path, "word/document.xml")


def test_xlsx_retains_sheet_cell_and_formula(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Budget"
    workbook.active.append(["Item", "Amount"])
    workbook.active.append(["Travel", 42])
    workbook.active["B3"] = "=SUM(B2:B2)"
    workbook.create_sheet("Notes")["A1"] = "Review annually"
    path = tmp_path / "budget.xlsx"
    workbook.save(path)
    docs = load_document(path)
    assert len(docs) == 2
    assert "B2: 42" in docs[0].content
    assert "=SUM(B2:B2)" in docs[0].content
    assert docs[0].metadata.section_title == "Budget"
    assert docs[1].metadata.page_number == 2


def test_pptx_retains_slide_text_table_and_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Annual policy review"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "HR"
    slide.notes_slide.notes_text_frame.text = "Approve by December"
    path = tmp_path / "review.pptx"
    presentation.save(path)
    doc = load_document(path)[0]
    assert "Annual policy review" in doc.content
    assert "Owner | HR" in doc.content
    assert "Approve by December" in doc.content
    assert doc.metadata.page_number == 1


def test_docx_does_not_duplicate_runs(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Unique policy sentence.")
    path = tmp_path / "policy.docx"
    document.save(path)
    result = load_document(path, {"source_file": "original.docx"})[0]
    assert result.content.count("Unique policy sentence.") == 1
    assert result.metadata.source_file == "original.docx"


@pytest.fixture
def service(tmp_path, monkeypatch):
    from backend.services.document_service import DocumentService
    from backend.retrieval.bm25 import BM25SearchIndex

    monkeypatch.setattr("backend.api.dependencies.get_telemetry_service", lambda: MagicMock())
    vectors = MagicMock()
    vectors._collection = None
    embeddings = MagicMock()
    embeddings.embed_chunks.side_effect = lambda texts: [[0.1, 0.2] for _ in texts]
    assets = MagicMock()
    assets.list_assets.return_value = []
    instance = DocumentService(
        vector_store=vectors,
        embedding_service=embeddings,
        bm25_index=BM25SearchIndex(storage_dir=str(tmp_path / "bm25")),
        image_asset_manager=assets,
        vision_cache_manager=MagicMock(),
        storage_dir=str(tmp_path / "uploads"),
    )
    yield instance
    instance._ingestion_executor.shutdown(wait=True)


def finish(service):
    service._ingestion_executor.submit(lambda: None).result(timeout=15)


def test_upload_indexes_and_failed_parser_never_becomes_ready(service):
    good = service.upload_document("policy.txt", b"Employees receive twenty days of annual paid leave.")
    finish(service)
    assert service.get_ingestion_status(good.document_id).status == "READY"
    bad = service.upload_document("bad.json", b'{"bad":')
    finish(service)
    status = service.get_ingestion_status(bad.document_id)
    assert status.status == "FAILED"
    assert status.can_retry and not status.text_ready


def test_failed_upload_can_be_deleted(service):
    upload = service.upload_document("bad.json", b'{"bad":')
    finish(service)
    result = service.delete_document(upload.document_id)
    assert result is not None
    assert result["status"] == "deleted"
    assert not list(service.storage_dir.glob(f"{upload.document_id}_*"))


def test_failed_upload_is_deduplicated_to_retryable_job(service):
    content = b'{"bad":'
    upload = service.upload_document("bad.json", content)
    finish(service)
    from backend.services.document_service import DuplicateDocumentError

    with pytest.raises(DuplicateDocumentError) as exc_info:
        service.upload_document("renamed.json", content)
    assert exc_info.value.document_id == upload.document_id


def test_incomplete_embeddings_fail_before_index_write(service):
    service.embedding_service.embed_chunks.side_effect = lambda texts: []
    upload = service.upload_document("policy.txt", b"Employees receive twenty days of annual paid leave.")
    finish(service)
    assert service.get_ingestion_status(upload.document_id).status == "FAILED"
    service.vector_store.add_chunks.assert_not_called()


def test_partial_index_failure_is_cleaned(service):
    service.bm25_index.save = MagicMock(side_effect=OSError("disk full"))
    upload = service.upload_document("policy.txt", b"Employees receive twenty days of annual paid leave.")
    finish(service)
    assert service.get_ingestion_status(upload.document_id).status == "FAILED"
    service.vector_store.delete_by_document_id.assert_called_once_with(upload.document_id)
    assert service.bm25_index.entries == []
    assert service.docstore == {}


def test_queue_failure_clears_duplicate_reservation(service):
    service._ingestion_executor.shutdown(wait=True)
    with pytest.raises(RuntimeError):
        service.upload_document("policy.txt", b"Annual leave policy")
    assert service._pending_hashes == {}
    assert service._ingestion_jobs == {}
    assert service._stored_files == {}
    assert not list(service.storage_dir.iterdir())


def test_ready_retry_is_idempotent(service):
    upload = service.upload_document("policy.txt", b"Employees receive twenty days of annual paid leave.")
    finish(service)
    assert service.retry_document(upload.document_id).status == "READY"
    finish(service)
    service.vector_store.add_chunks.assert_called_once()


@pytest.mark.parametrize("filename", ["bad.exe", "bad:stream.txt", "a" * 181 + ".txt"])
def test_invalid_upload_rejected_before_storage(service, filename):
    with pytest.raises(ValueError):
        service.upload_document(filename, b"content")
    assert not list(service.storage_dir.iterdir())
