"""Text and table extraction for modern Excel and PowerPoint files."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from backend.ingestion.loaders.base import BaseLoader
from backend.ingestion.loaders.validation import validate_office_archive
from backend.models.document import DocumentType, RawDocument


class SpreadsheetLoader(BaseLoader):
    def supports(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".xlsx"

    def load(self, file_path: Path, base_metadata: dict[str, Any] | None = None) -> list[RawDocument]:
        from openpyxl import load_workbook

        validate_office_archive(file_path, "xl/workbook.xml")
        meta = self._build_base_metadata(file_path, DocumentType.XLSX, base_metadata)
        # Preserve formulas explicitly; cached values can be absent or stale.
        workbook = load_workbook(file_path, read_only=True, data_only=False, keep_links=False)
        documents = []
        cells_seen = 0
        try:
            for index, sheet in enumerate(workbook.worksheets):
                if (sheet.max_row or 0) * (sheet.max_column or 0) > 1_000_000:
                    raise ValueError("Worksheet exceeds the one-million-cell limit. Split the workbook.")
                lines = []
                for row in sheet.iter_rows():
                    cells_seen += len(row)
                    if cells_seen > 1_000_000:
                        raise ValueError("Workbook exceeds the one-million-cell limit. Split the workbook.")
                    values = [f"{cell.coordinate}: {cell.value}" for cell in row if cell.value is not None]
                    if values:
                        lines.append("; ".join(values))
                if lines:
                    documents.append(
                        RawDocument(
                            content="\n".join(lines),
                            metadata=meta.model_copy(
                                update={
                                    "section_title": sheet.title,
                                    "section_path": sheet.title,
                                    "page_number": index + 1,
                                    "internal_page_index": index,
                                    "total_pages": len(workbook.worksheets),
                                    "extra": {**meta.extra, "sheet_name": sheet.title, "formulas": "unevaluated"},
                                }
                            ),
                        )
                    )
        finally:
            workbook.close()
        return documents


class PresentationLoader(BaseLoader):
    def supports(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".pptx"

    def load(self, file_path: Path, base_metadata: dict[str, Any] | None = None) -> list[RawDocument]:
        from pptx import Presentation

        validate_office_archive(file_path, "ppt/presentation.xml")
        meta = self._build_base_metadata(file_path, DocumentType.PPTX, base_metadata)
        presentation = Presentation(str(file_path))
        documents = []

        def texts(shapes):
            for shape in shapes:
                if shape.has_text_frame:
                    yield shape.text_frame.text
                if shape.has_table:
                    for row in shape.table.rows:
                        yield " | ".join(cell.text for cell in row.cells)
                if hasattr(shape, "shapes"):
                    yield from texts(shape.shapes)

        for index, slide in enumerate(presentation.slides):
            parts = [text for text in texts(slide.shapes) if text.strip()]
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes and notes.text.strip():
                    parts.append("Speaker notes: " + notes.text)
            if parts:
                documents.append(
                    RawDocument(
                        content="\n\n".join(parts),
                        metadata=meta.model_copy(
                            update={
                                "page_number": index + 1,
                                "internal_page_index": index,
                                "total_pages": len(presentation.slides),
                                "section_title": f"Slide {index + 1}",
                            }
                        ),
                    )
                )
        return documents
